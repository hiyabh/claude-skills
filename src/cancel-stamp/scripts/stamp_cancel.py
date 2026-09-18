# -*- coding: utf-8 -*-
"""Stamp a DOCX or PDF as cancelled/void, without touching a single character.

Draws a transparent full-page overlay - either a diagonal BAND (two parallel
lines with the notice written between them) or a CROSS (an X with the notice in
the gap at the crossing point) - and lays it over the text layer of every page.

Usage:
  python stamp_cancel.py "<file.docx|file.pdf>" [options]

  --style band|cross     default: band
  --text LINE            repeatable; first line is the headline (bigger)
  --opacity 0.0-1.0      default: 0.41 (faded, keeps the document readable)
  --color R,G,B          default: 200,16,16 (cancellation red)
  --pages all|first      default: all
  --out PATH             write a copy instead of stamping in place
  --font PATH            TTF used for the notice (must cover the script)
  --size-major / --size-minor   font size in px at 200 DPI

Why a raster overlay and not DrawingML/PDF shapes: Word does not honor z-order
between a shape and a behind-text image on PDF export, so a letterhead
background can swallow the stamp. An in-front-of-text transparent PNG always
wins, in Word, in print, and in the exported PDF alike.
"""
import argparse
import hashlib
import io
import math
import re
import shutil
import sys
import tempfile
import time
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont, features

DEFAULT_LINES = ["בוטל - לא נשלח"]
DEFAULT_COLOR = (200, 16, 16)
DEFAULT_OPACITY = 0.41
DEFAULT_FONT = r"C:\Windows\Fonts\arialbd.ttf"

DPI = 200                     # overlay raster resolution
SIZE_MAJOR_PX = 150           # headline size at DPI
SIZE_MINOR_PX = 78            # every following line
LINE_SPACING_PX = 30
LINE_WIDTH_PX = 16            # thickness of each stamp rule
EDGE_MARGIN_PX = 150          # keeps the stamp clear of the paper edges
TEXT_GAP_PX = 46              # clear space between the notice and each rule
TEXT_ALPHA_BOOST = 1.15       # the notice reads slightly stronger than the rules

EMU_PER_CM = 360000
CM_PER_INCH = 2.54
PT_PER_INCH = 72.0
OVERLAY_Z = "503316490"       # above the letterhead background AND company stamps
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
WD_ACTIVE_END_PAGE = 3

HAS_RAQM = features.check("raqm")
TEXT_KW = {"direction": "rtl", "language": "he"} if HAS_RAQM else {}
# Runs that must stay in logical order inside reversed RTL text (numbers, Latin).
# Spaces are allowed inside a run so "Draft v2" keeps its word order.
LTR_TOKEN = r"[0-9A-Za-z][0-9A-Za-z.,:/+\-]*"
LTR_RUN = re.compile(rf"{LTR_TOKEN}(?:[ \t]+{LTR_TOKEN})*")
MIRRORED = str.maketrans("()[]{}<>", ")(][}{><")


def shape(text):
    """Visual-order text for PIL builds without libraqm (no bidi engine).

    Reverses the string, then un-reverses embedded Latin/number runs and mirrors
    bracket glyphs - enough for one-line Hebrew notices with dates or codes in
    them. Not a full UAX#9 implementation; nested bidi levels are out of scope.
    """
    if HAS_RAQM:
        return text
    flipped = text[::-1].translate(MIRRORED)
    return LTR_RUN.sub(lambda m: m.group(0)[::-1], flipped)


def fitted_font(draw, text, px, max_width, font_path):
    while px > 10:
        font = ImageFont.truetype(font_path, px)
        if draw.textlength(text, font=font, **TEXT_KW) <= max_width:
            return font
        px -= 4
    return ImageFont.truetype(font_path, px)


def measure(cfg, max_width):
    """Fit every notice line to max_width; return fonts, heights, block height."""
    probe = ImageDraw.Draw(Image.new("RGBA", (1, 1)))
    fonts, heights = [], []
    for i, text in enumerate(cfg.text):
        px = cfg.size_major if i == 0 else cfg.size_minor
        font = fitted_font(probe, shape(text), px, max_width, cfg.font)
        fonts.append(font)
        heights.append(font.getbbox(shape(text))[3])
    block = sum(heights) + LINE_SPACING_PX * (len(cfg.text) - 1)
    return fonts, heights, block


def render_notice(draw, center_x, top, cfg, fonts, heights, color):
    y = top
    for text, font, height in zip(cfg.text, fonts, heights):
        draw.text((center_x, y), shape(text), font=font, fill=color,
                  anchor="ma", **TEXT_KW)
        y += height + LINE_SPACING_PX


def band_height(cfg, length):
    return LINE_WIDTH_PX * 2 + TEXT_GAP_PX * 2 + measure(cfg, length - 2 * EDGE_MARGIN_PX)[2]


def build_band(cfg, length, line_color, text_color):
    """Flat band: rule - notice - rule. Rotated onto the page by the caller."""
    fonts, heights, _ = measure(cfg, length - 2 * EDGE_MARGIN_PX)
    height = band_height(cfg, length)
    band = Image.new("RGBA", (length, height), (0, 0, 0, 0))
    draw = ImageDraw.Draw(band)
    half = LINE_WIDTH_PX // 2
    for y in (half, height - half - 1):
        draw.line([(0, y), (length - 1, y)], fill=line_color, width=LINE_WIDTH_PX)
    render_notice(draw, length / 2, LINE_WIDTH_PX + TEXT_GAP_PX,
                  cfg, fonts, heights, text_color)
    return band


def split_diagonal(p0, p1, band_top, band_bottom):
    """Split a diagonal into the two segments outside the horizontal notice band."""
    (x0, y0), (x1, y1) = p0, p1

    def at(y):
        return (x0 + (y - y0) / (y1 - y0) * (x1 - x0), y)

    lo, hi = (band_top, band_bottom) if y0 < y1 else (band_bottom, band_top)
    return [(p0, at(lo)), (at(hi), p1)]


def build_cross(cfg, w, h, line_color, text_color):
    """Two diagonals crossing the page, broken around a centered notice."""
    img = Image.new("RGBA", (w, h), (0, 0, 0, 0))
    draw = ImageDraw.Draw(img)
    fonts, heights, block = measure(cfg, w - 2 * EDGE_MARGIN_PX - 120)
    top = (h - block) // 2
    corners = [((EDGE_MARGIN_PX, EDGE_MARGIN_PX), (w - EDGE_MARGIN_PX, h - EDGE_MARGIN_PX)),
               ((w - EDGE_MARGIN_PX, EDGE_MARGIN_PX), (EDGE_MARGIN_PX, h - EDGE_MARGIN_PX))]
    for p0, p1 in corners:
        for seg in split_diagonal(p0, p1, top - TEXT_GAP_PX,
                                  top + block + TEXT_GAP_PX):
            draw.line([tuple(map(int, seg[0])), tuple(map(int, seg[1]))],
                      fill=line_color, width=LINE_WIDTH_PX)
    render_notice(draw, w / 2, top, cfg, fonts, heights, text_color)
    return img


def build_overlay(cfg, w, h):
    """Transparent PNG stream sized w x h pixels carrying the cancellation art."""
    line_color = cfg.color + (round(255 * cfg.opacity),)
    text_color = cfg.color + (round(255 * min(1.0, cfg.opacity * TEXT_ALPHA_BOOST)),)
    if cfg.style == "cross":
        img = build_cross(cfg, w, h, line_color, text_color)
    else:
        # Band on the page diagonal, shortened so BOTH rules end inside the
        # margins (a band of the full diagonal pushes its rules past the corners)
        img = Image.new("RGBA", (w, h), (0, 0, 0, 0))
        avail_w, avail_h = w - 2 * EDGE_MARGIN_PX, h - 2 * EDGE_MARGIN_PX
        rad = math.atan2(avail_h, avail_w)
        cos_a, sin_a = math.cos(rad), math.sin(rad)
        thickness = band_height(cfg, int(math.hypot(avail_w, avail_h)))
        length = int(min((avail_w - thickness * sin_a) / cos_a,
                         (avail_h - thickness * cos_a) / sin_a))
        band = build_band(cfg, length, line_color, text_color)
        rotated = band.rotate(math.degrees(rad), expand=True, resample=Image.BICUBIC)
        img.alpha_composite(rotated, ((w - rotated.width) // 2,
                                      (h - rotated.height) // 2))
    out = io.BytesIO()
    img.save(out, format="PNG")
    out.seek(0)
    return out


def anchor_xml(rid, cx, cy, shape_id):
    from docx.oxml.ns import nsdecls
    return (
        f'<w:drawing {nsdecls("w", "wp", "a", "pic", "r")}>'
        f'<wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0"'
        f' relativeHeight="{OVERLAY_Z}" behindDoc="0" locked="0"'
        f' layoutInCell="0" allowOverlap="1">'
        f'<wp:simplePos x="0" y="0"/>'
        f'<wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH>'
        f'<wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>'
        f'<wp:extent cx="{cx}" cy="{cy}"/>'
        f'<wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/>'
        f'<wp:docPr id="{shape_id}" name="CancelStamp{shape_id}"/>'
        f'<wp:cNvGraphicFramePr/>'
        f'<a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">'
        f'<pic:pic><pic:nvPicPr><pic:cNvPr id="{shape_id}" name="cancel.png"/>'
        f'<pic:cNvPicPr/></pic:nvPicPr>'
        f'<pic:blipFill><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></pic:blipFill>'
        f'<pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
        f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr>'
        f'</pic:pic></a:graphicData></a:graphic></wp:anchor></w:drawing>'
    )


def text_sha(doc):
    joined = "\n".join(p.text for p in doc.paragraphs)
    return hashlib.sha1(joined.encode("utf-8")).hexdigest()


def first_paragraph_of_each_page(path):
    """Pagination via Word COM, so page 2+ gets its own overlay copy."""
    try:
        import win32com.client
    except ImportError:
        print("note: pywin32 missing - stamping page 1 only")
        return [0]
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    try:
        doc = word.Documents.Open(str(path), ReadOnly=True,
                                  ConfirmConversions=False, AddToRecentFiles=False)
        first = {}
        for k in range(1, doc.Paragraphs.Count + 1):
            rng = doc.Paragraphs(k).Range
            rng.Collapse(1)
            first.setdefault(rng.Information(WD_ACTIVE_END_PAGE), k - 1)
        doc.Close(False)
    finally:
        word.Quit()
    return [first[pg] for pg in sorted(first)]


def px_for_emu(emu):
    return int(emu / EMU_PER_CM / CM_PER_INCH * DPI)


def stamp_docx(cfg, path):
    from docx import Document
    from docx.oxml import parse_xml

    targets = first_paragraph_of_each_page(path) if cfg.pages == "all" else [0]
    doc = Document(str(path))
    before = text_sha(doc)
    section = doc.sections[0]
    cx, cy = int(section.page_width), int(section.page_height)
    png = build_overlay(cfg, px_for_emu(cx), px_for_emu(cy))
    rid, _ = doc.part.get_or_add_image(png)
    for n, idx in enumerate(targets, start=1):
        run = doc.paragraphs[idx].add_run()
        run._element.append(parse_xml(anchor_xml(rid, cx, cy, 9000 + n)))
        print(f"page {n}: cancel stamp anchored to paragraph {idx}")
    if text_sha(doc) != before:
        raise RuntimeError("text changed - aborting without save")
    doc.save(str(path))


def stamp_pdf(cfg, path):
    import fitz
    doc = fitz.open(str(path))
    pages = doc if cfg.pages == "all" else [doc[0]]
    cache = {}
    for page in pages:
        key = (round(page.rect.width), round(page.rect.height))
        if key not in cache:
            w = int(key[0] / PT_PER_INCH * DPI)
            h = int(key[1] / PT_PER_INCH * DPI)
            cache[key] = build_overlay(cfg, w, h).getvalue()
        page.insert_image(page.rect, stream=cache[key], overlay=True)
        print(f"page {page.number + 1}: cancel stamp drawn")
    # Full rewrite with deflate: an incremental save stores the overlay raster
    # uncompressed and inflates the file by ~15 MB per page.
    tmp = path.with_suffix(path.suffix + ".stamped")
    doc.save(str(tmp), deflate=True, deflate_images=True, garbage=3)
    doc.close()
    tmp.replace(path)


def parse_args(argv):
    ap = argparse.ArgumentParser(description="Stamp a DOCX/PDF as cancelled.")
    ap.add_argument("file")
    ap.add_argument("--style", choices=("band", "cross"), default="band")
    ap.add_argument("--text", action="append", default=None,
                    help="notice line; repeat for more lines (first = headline)")
    ap.add_argument("--opacity", type=float, default=DEFAULT_OPACITY)
    ap.add_argument("--color", default=",".join(map(str, DEFAULT_COLOR)))
    ap.add_argument("--pages", choices=("all", "first"), default="all")
    ap.add_argument("--out")
    ap.add_argument("--font", default=DEFAULT_FONT)
    ap.add_argument("--size-major", type=int, default=SIZE_MAJOR_PX)
    ap.add_argument("--size-minor", type=int, default=SIZE_MINOR_PX)
    cfg = ap.parse_args(argv)
    cfg.text = cfg.text or list(DEFAULT_LINES)
    cfg.color = tuple(int(v) for v in cfg.color.split(","))
    if not 0.0 < cfg.opacity <= 1.0:
        ap.error("--opacity must be within (0, 1]")
    if len(cfg.color) != 3:
        ap.error("--color must be R,G,B")
    return cfg


def main(argv):
    cfg = parse_args(argv)
    src = Path(cfg.file).resolve()
    target = Path(cfg.out).resolve() if cfg.out else src
    if target != src:
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(src, target)
    else:
        backup = Path(tempfile.gettempdir()) / f"{src.stem}.{int(time.time())}.pre-cancel{src.suffix}"
        shutil.copy2(src, backup)
        print("backup:", backup)

    suffix = target.suffix.lower()
    if suffix == ".docx":
        stamp_docx(cfg, target)
    elif suffix == ".pdf":
        stamp_pdf(cfg, target)
    else:
        raise SystemExit(f"unsupported file type: {suffix} (use .docx or .pdf)")
    print(f"OK {target.name} ({target.stat().st_size} bytes)")


if __name__ == "__main__":
    main(sys.argv[1:])
