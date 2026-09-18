"""Build RTL Hebrew Excel summary from Gemini OCR JSON results.

Usage:
    python build_xlsx.py <gemini_json_file> [<another_json>...] --out summary.xlsx

Reads one or more Gemini OCR result JSONs (different prompt types) and merges
into a single RTL Excel file with the standard Hebrew column schema.
"""
import json, os, sys, argparse
from openpyxl import Workbook

# Windows pipes stdout/stderr through the ANSI code page by default; force
# UTF-8 so Hebrew paths/values in progress output don't crash the run.
for _stream in (sys.stdout, sys.stderr):
    if hasattr(_stream, 'reconfigure'):
        _stream.reconfigure(encoding='utf-8', errors='replace')
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

UNCLEAR = "לא ברור"

COLUMNS = [
    "סמל מוטב", "שם המטפל", "טלפון", "שם המורה", "ת.ז.",
    "חודש התלוש", "סכום משרד החינוך למוסד (₪)",
    "סכום מענק בתלוש (₪)", "הפרשה לגמל 5% (₪)",
    "מקור", "הערות",
]


def fmt_months(months):
    if not months:
        return UNCLEAR
    return ", ".join(f'{m.get("חודש","?")}={m.get("סכום","?")}' for m in months if m)


def rows_from_bitsua(data, src_file):
    """Convert one bitsua-type result to row dicts."""
    if 'error' in data:
        return []
    inst = data.get('שם_המוסד') or UNCLEAR
    sml = data.get('סמל_מוטב') or UNCLEAR
    addr = data.get('כתובת_המוסד') or ''
    phone = data.get('טלפון_איש_קשר') or UNCLEAR
    contact = data.get('שם_איש_קשר') or ''
    period = data.get('תקופת_דיווח') or ''
    notes_base = " | ".join(filter(None, [
        f'תקופת דיווח: {period}' if period else None,
        f'כתובת המוסד: {addr}' if addr else None,
        f'איש קשר: {contact}' if contact else None,
    ]))
    out = []
    for t in data.get('מורים', []) or []:
        tid = t.get('ת.ז') or UNCLEAR
        last = t.get('שם_משפחה') or ''
        first = t.get('שם_פרטי') or ''
        full_name = f'{first} {last}'.strip() or UNCLEAR
        phone_t = t.get('טלפון_זמין') or phone
        moe = t.get('סכום_משרד_החינוך')
        paid = t.get('סהכ_שולם_למורה')
        months = t.get('תשלומים_חודשיים', [])
        # 5% from amount actually paid (per Hebrew Geulah convention)
        base = paid if paid else moe
        gemel = round(base * 0.05, 2) if base else UNCLEAR
        out.append({
            "סמל מוטב": sml, "שם המטפל": inst, "טלפון": phone_t,
            "שם המורה": full_name, "ת.ז.": tid,
            "חודש התלוש": fmt_months(months),
            "סכום משרד החינוך למוסד (₪)": moe if moe else UNCLEAR,
            "סכום מענק בתלוש (₪)": paid if paid else UNCLEAR,
            "הפרשה לגמל 5% (₪)": gemel,
            "מקור": src_file,
            "הערות": notes_base,
        })
    return out


def rows_from_generic(data, src_key):
    """Convert one generic-type result to row dicts.
    Type A or B → single teacher row. Type C → table rows. D/E → none."""
    if 'error' in data:
        return []
    dtype = data.get('סוג_מסמך', '')
    sml = data.get('סמל_מוטב') or UNCLEAR
    inst = data.get('שם_מוסד') or UNCLEAR
    amuta = data.get('מס_עמותה') or ''
    full_inst = f'{inst} (עמותה {amuta})' if amuta else inst

    if dtype in ('A', 'B'):
        t = data.get('פרטי_מורה') or {}
        if not t:
            return []
        tid = t.get('ת.ז') or UNCLEAR
        full_name = t.get('שם_מלא') or UNCLEAR
        phone = t.get('טלפון') or UNCLEAR
        addr = t.get('כתובת') or ''
        month = t.get('חודש_תלוש') or UNCLEAR
        grant = t.get('סכום_מענק')
        salary = t.get('סכום_שכר_כולל')
        gemel_paid = t.get('הפרשה_לגמל')
        gemel = round(grant * 0.05, 2) if grant else UNCLEAR
        notes = []
        if dtype == 'A':
            notes.append('סוג: אישור משרד החינוך')
        else:
            notes.append('סוג: תלוש שכר')
        if addr: notes.append(f'כתובת: {addr}')
        if salary: notes.append(f'שכר חודשי כולל: {salary}')
        if gemel_paid: notes.append(f'הפרשה בפועל בתלוש: {gemel_paid}')
        return [{
            "סמל מוטב": sml, "שם המטפל": full_inst, "טלפון": phone,
            "שם המורה": full_name, "ת.ז.": tid,
            "חודש התלוש": month,
            "סכום משרד החינוך למוסד (₪)": grant if grant else UNCLEAR,
            "סכום מענק בתלוש (₪)": grant if grant else UNCLEAR,
            "הפרשה לגמל 5% (₪)": gemel,
            "מקור": src_key,
            "הערות": " | ".join(notes),
        }]
    elif dtype == 'C':
        out = []
        for t in data.get('מורים_בטבלה', []) or []:
            tid = t.get('ת.ז') or UNCLEAR
            last = t.get('שם_משפחה') or ''
            first = t.get('שם_פרטי') or ''
            full_name = f'{first} {last}'.strip() or UNCLEAR
            moe = t.get('סכום_משרד_החינוך')
            paid = t.get('סהכ_שולם_למורה')
            months_desc = t.get('חודשים', '')
            base = paid if paid else moe
            gemel = round(base * 0.05, 2) if base else UNCLEAR
            out.append({
                "סמל מוטב": sml, "שם המטפל": full_inst, "טלפון": UNCLEAR,
                "שם המורה": full_name, "ת.ז.": tid,
                "חודש התלוש": months_desc or UNCLEAR,
                "סכום משרד החינוך למוסד (₪)": moe if moe else UNCLEAR,
                "סכום מענק בתלוש (₪)": paid if paid else UNCLEAR,
                "הפרשה לגמל 5% (₪)": gemel,
                "מקור": src_key,
                "הערות": "",
            })
        return out
    return []


def write_xlsx(rows, out_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "סיכום"
    ws.sheet_view.rightToLeft = True

    header_font = Font(name="David", size=12, bold=True, color="000000")
    header_fill = PatternFill("solid", fgColor="D9D9D9")
    thin = Side(border_style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center_wrap = Alignment(horizontal="center", vertical="center", wrap_text=True, readingOrder=2)
    right_wrap = Alignment(horizontal="right", vertical="center", wrap_text=True, readingOrder=2)

    for col_idx, h in enumerate(COLUMNS, 1):
        c = ws.cell(row=1, column=col_idx, value=h)
        c.font = header_font; c.fill = header_fill; c.border = border
        c.alignment = center_wrap

    data_font = Font(name="David", size=11, color="000000")
    for r_idx, row in enumerate(rows, 2):
        for c_idx, h in enumerate(COLUMNS, 1):
            v = row.get(h, "")
            c = ws.cell(row=r_idx, column=c_idx, value=v)
            c.font = data_font; c.border = border
            if isinstance(v, (int, float)):
                c.alignment = Alignment(horizontal="right", vertical="center", readingOrder=2)
                c.number_format = "#,##0.00"
            else:
                c.alignment = right_wrap

    widths = [12, 36, 16, 24, 14, 28, 22, 22, 18, 36, 60]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    wb.save(out_path)


def verify(rows):
    """Sanity check: 5% × paid should equal gemel."""
    err = 0
    total = 0
    n_clear = 0
    n_unclear_cells = 0
    for row in rows:
        grant = row.get("סכום מענק בתלוש (₪)")
        gemel = row.get("הפרשה לגמל 5% (₪)")
        if isinstance(grant, (int, float)) and isinstance(gemel, (int, float)):
            if abs(gemel - round(grant * 0.05, 2)) > 0.05:
                err += 1
            total += grant
            n_clear += 1
        for v in row.values():
            if isinstance(v, str) and v == UNCLEAR:
                n_unclear_cells += 1
    return {"rows": len(rows), "clear_grant": n_clear, "sum": total,
            "calc_errors": err, "unclear_cells": n_unclear_cells}


def main():
    p = argparse.ArgumentParser()
    p.add_argument('json_files', nargs='+')
    p.add_argument('--out', required=True)
    p.add_argument('--type', default='auto', choices=['auto', 'bitsua', 'generic'])
    args = p.parse_args()

    all_rows = []
    for jf in args.json_files:
        data = json.load(open(jf, encoding='utf-8'))
        # auto-detect: bitsua-type top-level keys map to instruments;
        # generic-type top-level keys are file/page identifiers
        for key, value in data.items():
            if isinstance(value, dict):
                # Detect: bitsua has "מורים" at top, generic has "סוג_מסמך"
                if 'מורים' in value and 'מורים_בטבלה' not in value:
                    # bitsua format: each top key is an institution
                    all_rows.extend(rows_from_bitsua(value, key))
                elif 'סוג_מסמך' in value:
                    all_rows.extend(rows_from_generic(value, key))

    write_xlsx(all_rows, args.out)
    stats = verify(all_rows)
    print(f'Saved: {args.out}')
    for k, v in stats.items():
        print(f'  {k}: {v}')


if __name__ == '__main__':
    main()
