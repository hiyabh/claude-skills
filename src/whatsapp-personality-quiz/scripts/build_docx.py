"""
Build an RTL Hebrew DOCX from an edited shiur Markdown file.

Usage:
    python build_docx.py "<path-to-md>"

Produces a .docx with the same basename next to the .md.
RTL full, David font, black & white, A4, page numbers.

Layout (shiur-editor house style):
  - Body / psukim / lists: 14pt, justified (both edges).
  - Headings: H1 title 16pt centered; H2-H4 14pt justified; all bold.
  - Font stays David (IRON RULE) — do NOT drop it.
Pasuk citations (markdown blockquotes `> `) are rendered indented + bold.

These sizes/alignments are overridden locally here, not in the vendored
hebrew_docx_template.py sitting next to this file (that copy is the shared
helper library and stays generic).
"""
import sys
import os
import re

# Claude Code captures stdout through a pipe. On Windows that makes Python fall
# back to the ANSI code page, so printing a Hebrew filename dies with
# UnicodeEncodeError instead of reporting success. Force UTF-8 before we print.
for _stream in (sys.stdout, sys.stderr):
    if hasattr(_stream, 'reconfigure'):
        _stream.reconfigure(encoding='utf-8', errors='replace')

# Template resolution, portable and self-contained. The vendored copy shipping
# beside this script WINS: its directory goes first, so an installed bundle
# needs no external files. ~/.claude/templates is appended AFTER it, so a
# machine that happens to have the shared global copy keeps working — it is a
# fallback, never an override.
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
if SCRIPT_DIR not in sys.path:
    sys.path.insert(0, SCRIPT_DIR)
_GLOBAL_TEMPLATES = os.path.expanduser('~/.claude/templates')
if os.path.isdir(_GLOBAL_TEMPLATES) and _GLOBAL_TEMPLATES not in sys.path:
    sys.path.append(_GLOBAL_TEMPLATES)

# docx is checked before the template: the template imports docx itself, so a
# missing python-docx would otherwise surface as the wrong (template) error.
try:
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
except ImportError:
    print('ERROR: חסרה הספרייה python-docx — בלעדיה אי אפשר לבנות את ה-DOCX.')
    print(f'  התקנה: "{sys.executable}" -m pip install python-docx')
    sys.exit(1)

try:
    from hebrew_docx_template import (
        set_rtl_paragraph, style_run, add_page_numbers, set_run_lang_hebrew,
        HEB_FONT, BLACK,
    )
except ImportError:
    print('ERROR: לא נמצא hebrew_docx_template.py.')
    print(f'  הוא אמור לשבת ליד הסקריפט הזה, בתיקייה: {SCRIPT_DIR}')
    print('  ההתקנה חלקית — התקן את הסקיל מחדש מהחבילה המלאה.')
    sys.exit(1)

# House style sizes (pt) — see module docstring.
BODY_SIZE = 14
QUOTE_SIZE = 14
LIST_SIZE = 14
TITLE_LEVEL = 1
HEADING_SIZES = {1: 16, 2: 14, 3: 14, 4: 14}


def add_heading_local(doc, text, level=1):
    """H1 title centered; deeper headings justified. All bold, David."""
    p = doc.add_paragraph()
    p.alignment = (WD_ALIGN_PARAGRAPH.CENTER if level == TITLE_LEVEL
                   else WD_ALIGN_PARAGRAPH.JUSTIFY)
    set_rtl_paragraph(p)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    style_run(run, size=HEADING_SIZES.get(level, BODY_SIZE), bold=True)
    return p


def add_body_para(doc, text, size=BODY_SIZE):
    """Justified body paragraph, preserving **bold** spans."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_rtl_paragraph(p)
    p.paragraph_format.space_after = Pt(4)
    for part in re.split(r'(\*\*.+?\*\*)', text):
        if part.startswith('**') and part.endswith('**'):
            run = p.add_run(part[2:-2])
            style_run(run, size=size, bold=True)
        elif part:
            run = p.add_run(part)
            style_run(run, size=size)
    return p


def add_hyperlink_run(paragraph, url, display_text, size=11):
    """Real clickable OOXML hyperlink run (not just visible text). Kept black +
    underlined rather than the usual blue, per the B&W-only Iron Rule — the
    underline signals "clickable" without breaking the house color rule."""
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)
    hyperlink.set(qn('w:history'), '1')

    r = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:cs'), HEB_FONT)
    rFonts.set(qn('w:ascii'), HEB_FONT)
    rFonts.set(qn('w:hAnsi'), HEB_FONT)
    rPr.append(rFonts)
    u = OxmlElement('w:u')
    u.set(qn('w:val'), 'single')
    rPr.append(u)
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '000000')
    rPr.append(color)
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), str(int(size * 2)))
    rPr.append(sz)
    szCs = OxmlElement('w:szCs')
    szCs.set(qn('w:val'), str(int(size * 2)))
    rPr.append(szCs)
    rPr.append(OxmlElement('w:rtl'))
    set_run_lang_hebrew(rPr)
    r.append(rPr)

    t = OxmlElement('w:t')
    t.text = display_text
    t.set(qn('xml:space'), 'preserve')
    r.append(t)

    hyperlink.append(r)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_source_line(doc, text):
    """Source/subtitle line right under the H1 (e.g. 'מקור: <youtube url>') —
    centered, smaller, the label as plain text and the URL as a real clickable
    hyperlink, distinct from body paragraphs."""
    label, sep, url = text.partition(': ')
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_rtl_paragraph(p)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(10)
    if sep and url:
        prefix_run = p.add_run(label + sep)
        style_run(prefix_run, size=11, bold=False)
        add_hyperlink_run(p, url, url, size=11)
    else:
        run = p.add_run(text)
        style_run(run, size=11, bold=False)
    return p


def add_quote_para(doc, text):
    """Pasuk citation — indented both sides, bold, justified, slight spacing."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_rtl_paragraph(p)
    p.paragraph_format.right_indent = Cm(1.0)
    p.paragraph_format.left_indent = Cm(1.0)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    style_run(run, size=QUOTE_SIZE, bold=True)
    return p


def add_list_item(doc, content, marker):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_rtl_paragraph(p)
    # In a bidi paragraph w:left is the start (right) edge, so the hanging
    # marker must indent from left_indent. Using right_indent pushes the
    # bullet past the right page margin.
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.right_indent = 0
    p.paragraph_format.first_line_indent = Cm(-0.5)
    p.paragraph_format.space_after = Pt(2)
    rb = p.add_run(marker)
    style_run(rb, size=LIST_SIZE, bold=True)
    r = p.add_run(content)
    style_run(r, size=LIST_SIZE)


def setup_document():
    doc = Document()
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    sectPr = section._sectPr
    sectPr.append(OxmlElement('w:bidi'))
    sectPr.append(OxmlElement('w:rtlGutter'))
    add_page_numbers(section)

    style = doc.styles['Normal']
    style.font.name = HEB_FONT
    style.font.size = Pt(11)
    style.font.color.rgb = BLACK
    style_pPr = style.element.find(qn('w:pPr'))
    if style_pPr is None:
        style_pPr = OxmlElement('w:pPr')
        style.element.append(style_pPr)
    style_pPr.append(OxmlElement('w:bidi'))
    style_rPr = style.element.find(qn('w:rPr'))
    if style_rPr is None:
        style_rPr = OxmlElement('w:rPr')
        style.element.append(style_rPr)
    style_rPr.append(OxmlElement('w:rtl'))
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:cs'), HEB_FONT)
    rFonts.set(qn('w:ascii'), HEB_FONT)
    rFonts.set(qn('w:hAnsi'), HEB_FONT)
    style_rPr.append(rFonts)
    set_run_lang_hebrew(style_rPr)
    return doc


def render(doc, text):
    text = text.replace('—', '-').replace('–', '-')  # em/en dash -> hyphen
    for raw in text.split('\n'):
        s = raw.strip()
        if not s:
            continue
        if s.startswith('# '):
            add_heading_local(doc, s[2:].strip(), level=1); continue
        if s.startswith('#### '):
            add_heading_local(doc, s[5:].strip(), level=4); continue
        if s.startswith('### '):
            add_heading_local(doc, s[4:].strip(), level=3); continue
        if s.startswith('## '):
            add_heading_local(doc, s[3:].strip(), level=2); continue
        if s == '---':
            continue
        if s.startswith('מקור: '):
            add_source_line(doc, s); continue
        if s.startswith('> '):
            add_quote_para(doc, s[2:].strip()); continue
        if s.startswith('- ') or s.startswith('* '):
            content = re.sub(r'\*\*(.+?)\*\*', r'\1', s[2:])
            add_list_item(doc, content, '• '); continue
        m = re.match(r'^(\d+)\.\s+(.+)$', s)
        if m:
            content = re.sub(r'\*\*(.+?)\*\*', r'\1', m.group(2))
            add_list_item(doc, content, f"\u200F{m.group(1)}.\u200F ")
            continue
        add_body_para(doc, s, size=BODY_SIZE)


def main():
    if len(sys.argv) < 2:
        print('Usage: python build_docx.py "<path-to-md>"')
        sys.exit(1)
    md_path = sys.argv[1]
    if not os.path.isfile(md_path):
        print(f'ERROR: file not found: {md_path}')
        sys.exit(1)
    out_path = os.path.splitext(md_path)[0] + '.docx'
    with open(md_path, 'r', encoding='utf-8') as f:
        text = f.read()
    doc = setup_document()
    render(doc, text)
    doc.save(out_path)
    print(f'saved: {out_path}')
    print(f'size: {os.path.getsize(out_path)} bytes')


if __name__ == '__main__':
    main()
