"""
Shared RTL-Hebrew DOCX helpers - black & white only, David font.

Vendored copy: this file ships inside the shiur-editor skill so the bundle is
self-contained and needs nothing from outside its own directory. It is a helper
LIBRARY only - it is imported, never executed. build_docx.py owns the CLI.
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import re


MIN_PT = 12   # IRON RULE: no type smaller than 12pt anywhere in the document.


def _enforce_min(size):
    return max(size, MIN_PT)

BLACK = RGBColor(0, 0, 0)
HEB_FONT = 'David'

# ---------------------------------------------------------------------------
# TWO RTL RULES THAT MUST NOT BE BROKEN (verified against Word's own renderer,
# 2026-07-21). Breaking either one silently produces LEFT-aligned Hebrew:
#
# 1. ORDER: OOXML mandates a fixed child order inside <w:pPr>. w:bidi must come
#    BEFORE w:jc. Appending it afterwards makes Word drop the tag entirely and
#    the paragraph falls back to LTR. Always insert via insert_element_before.
#
# 2. NEVER set w:jc="right" on a Hebrew paragraph. Inside a bidi paragraph jc is
#    LOGICAL, not physical: "right" means end-of-line, which in RTL is the LEFT
#    side. Omit jc entirely — the line then starts at its natural start, i.e.
#    flush against the RIGHT margin. Only "center" is safe (it is symmetric).
# ---------------------------------------------------------------------------

# ECMA-376 CT_PPr child order.
_PPR_SEQ = (
    'w:pStyle', 'w:keepNext', 'w:keepLines', 'w:pageBreakBefore', 'w:framePr',
    'w:widowControl', 'w:numPr', 'w:suppressLineNumbers', 'w:pBdr', 'w:shd',
    'w:tabs', 'w:suppressAutoHyphens', 'w:kinsoku', 'w:wordWrap',
    'w:overflowPunct', 'w:topLinePunct', 'w:autoSpaceDE', 'w:autoSpaceDN',
    'w:bidi', 'w:adjustRightInd', 'w:snapToGrid', 'w:spacing', 'w:ind',
    'w:contextualSpacing', 'w:mirrorIndents', 'w:suppressOverlap', 'w:jc',
    'w:textDirection', 'w:textAlignment', 'w:textboxTightWrap', 'w:outlineLvl',
    'w:divId', 'w:cnfStyle', 'w:rPr', 'w:sectPr', 'w:pPrChange',
)


def successors_of(tag):
    return _PPR_SEQ[_PPR_SEQ.index(tag) + 1:]


def set_rtl_paragraph(p):
    """Mark a paragraph RTL. See RULE 1 — bidi is inserted, never appended."""
    pPr = p._p.get_or_add_pPr()
    if pPr.find(qn('w:bidi')) is None:
        bidi = OxmlElement('w:bidi')
        bidi.set(qn('w:val'), '1')
        pPr.insert_element_before(bidi, *successors_of('w:bidi'))
    rPr = pPr.find(qn('w:rPr'))
    if rPr is None:
        rPr = OxmlElement('w:rPr')
        pPr.insert_element_before(rPr, *successors_of('w:rPr'))
    if rPr.find(qn('w:rtl')) is None:
        rPr.append(OxmlElement('w:rtl'))


def apply_align(p, align):
    """Set alignment the RTL-safe way. See RULE 2 — RIGHT must stay unset."""
    if align == WD_ALIGN_PARAGRAPH.CENTER:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == WD_ALIGN_PARAGRAPH.JUSTIFY:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    # RIGHT (and anything else) → leave w:jc absent so RTL start = right margin

def set_rtl_run(run):
    """Force run to be RTL — required for proper number/bullet placement in mixed text."""
    rPr = run._r.get_or_add_rPr()
    rtl = OxmlElement('w:rtl')
    rPr.append(rtl)

def style_run(run, size=12, bold=False, italic=False, font=HEB_FONT):
    size = _enforce_min(size)
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = BLACK
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:cs'), font)
    rFonts.set(qn('w:ascii'), font)
    rFonts.set(qn('w:hAnsi'), font)
    # complex script size
    szCs = OxmlElement('w:szCs')
    szCs.set(qn('w:val'), str(int(size*2)))
    rPr.append(szCs)
    if bold:
        bCs = OxmlElement('w:bCs')
        rPr.append(bCs)
    # always mark as RTL run
    rtl = OxmlElement('w:rtl')
    rPr.append(rtl)
    set_run_lang_hebrew(rPr)

def set_run_lang_hebrew(rPr):
    """Tag run language as he-IL — without this, python-docx's default template
    proofs complex-script text as ar-SA and Word flags every Hebrew word."""
    lang = rPr.find(qn('w:lang'))
    if lang is None:
        lang = OxmlElement('w:lang')
        rPr.append(lang)
    lang.set(qn('w:val'), 'he-IL')
    lang.set(qn('w:bidi'), 'he-IL')

def add_para(doc, text, size=12, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=4):
    p = doc.add_paragraph()
    set_rtl_paragraph(p)
    apply_align(p, align)
    p.paragraph_format.space_after = Pt(space_after)
    if text:
        run = p.add_run(text)
        style_run(run, size=size, bold=bold, italic=italic)
    return p

def add_heading(doc, text, level=1):
    sizes = {1: 18, 2: 16, 3: 14, 4: 13}
    p = doc.add_paragraph()
    set_rtl_paragraph(p)
    apply_align(p, WD_ALIGN_PARAGRAPH.CENTER)  # IRON RULE: headings centered
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    style_run(run, size=sizes.get(level, 13), bold=True)
    return p


def add_page_numbers(section):
    """IRON RULE: every document gets centered page numbers in the footer (PAGE field)."""
    footer = section.footer
    footer.is_linked_to_previous = False
    para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    para.text = ''
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run()
    f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = 'PAGE'
    f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'end')
    run._r.append(f1); run._r.append(instr); run._r.append(f2)
    run.font.name = HEB_FONT; run.font.color.rgb = BLACK
    rPr = run._r.get_or_add_rPr()
    rf = OxmlElement('w:rFonts'); rf.set(qn('w:cs'), HEB_FONT); rf.set(qn('w:ascii'), HEB_FONT); rf.set(qn('w:hAnsi'), HEB_FONT)
    rPr.append(rf)
    set_run_lang_hebrew(rPr)

def set_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblBorders = OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        b = OxmlElement(f'w:{edge}')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), '4')
        b.set(qn('w:color'), '000000')
        tblBorders.append(b)
    tblPr.append(tblBorders)
    # bidi visual
    bidiV = OxmlElement('w:bidiVisual')
    tblPr.append(bidiV)

def shade_cell(cell, color='D9D9D9'):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color)
    tcPr.append(shd)

def fill_cell(cell, text, bold=False, size=12, header=False):
    cell.text = ''
    p = cell.paragraphs[0]
    set_rtl_paragraph(p)
    run = p.add_run(text)
    style_run(run, size=size, bold=bold or header)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    if header:
        shade_cell(cell, 'D9D9D9')  # light grey only

USABLE_CM = 17.0     # A4 (21 cm) minus 2 cm margins on each side
CHAR_CM = 0.21       # approx. width of one David 12pt Hebrew glyph
CELL_PAD_CM = 0.4    # left + right cell padding
HARD_MIN_CM = 1.2


def auto_col_widths(header, rows):
    """Proportional widths: a column holding long prose gets more of the 17 cm.

    Equal widths crush a 'problem and deliverable' column next to a '#' column,
    so weight each column by its mean cell length. The weight is square-rooted:
    raw means differ by 50x here, and a linear split starves the short columns
    until words break mid-syllable. Each column then gets a floor wide enough
    for its own longest word, so no heading is ever hyphenated by the table.
    """
    n = len(header)
    weights, floors = [], []
    for i in range(n):
        cells = [header[i]] + [r[i] for r in rows if i < len(r)]
        mean_len = sum(len(c) for c in cells) / len(cells)
        weights.append(max(2.0, mean_len) ** 0.5)
        longest_word = max((len(w) for c in cells for w in c.split()), default=1)
        floors.append(min(4.5, max(HARD_MIN_CM, longest_word * CHAR_CM + CELL_PAD_CM)))

    total = sum(weights)
    widths = [max(floors[i], USABLE_CM * weights[i] / total) for i in range(n)]
    # Give the overflow back proportionally, from the columns above their floor.
    overflow = sum(widths) - USABLE_CM
    while overflow > 0.01:
        slack = [max(0.0, widths[i] - floors[i]) for i in range(n)]
        if sum(slack) < 0.01:
            break
        take = overflow / sum(slack)
        widths = [widths[i] - slack[i] * take for i in range(n)]
        overflow = sum(widths) - USABLE_CM
    return widths


def _fix_layout(table, widths_cm):
    """Pin the column widths.

    Cell widths alone are advisory - Word and LibreOffice both re-flow the table
    unless the layout is declared fixed AND w:tblGrid carries the same widths.
    Without this a prose column gets crushed to the width of a '#' column.
    """
    tblPr = table._tbl.tblPr
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    tblPr.append(layout)
    grid = table._tbl.find(qn('w:tblGrid'))
    for col, w in zip(grid.findall(qn('w:gridCol')), widths_cm):
        col.set(qn('w:w'), str(int(w * 567)))   # cm -> twips


def _repeat_header(table):
    """Mark row 0 as a header row so it repeats on every page the table spans."""
    trPr = table.rows[0]._tr.get_or_add_trPr()
    trPr.append(OxmlElement('w:tblHeader'))


def add_table_md(doc, header, rows, col_widths=None):
    if col_widths is None:
        col_widths = auto_col_widths(header, rows)
    table = doc.add_table(rows=1+len(rows), cols=len(header))
    table.autofit = False
    set_table_borders(table)
    # header
    for i, h in enumerate(header):
        fill_cell(table.rows[0].cells[i], h, header=True, size=12)
    # rows
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            fill_cell(table.rows[r].cells[c], val, size=12)
    for i, w in enumerate(col_widths):
        for row in table.rows:
            row.cells[i].width = Cm(w)
    _fix_layout(table, col_widths)
    _repeat_header(table)
    return table

def parse_md_table(lines, start):
    """Parse md table starting at lines[start] which begins with '|'. Return (header, rows, end_index)."""
    header_line = lines[start].strip()
    if not header_line.startswith('|'): return None
    sep_line = lines[start+1].strip() if start+1 < len(lines) else ''
    if not re.match(r'^\|[\s\-:|]+\|$', sep_line): return None
    header = [c.strip() for c in header_line.strip('|').split('|')]
    rows = []
    i = start + 2
    while i < len(lines) and lines[i].strip().startswith('|'):
        cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
        # md bold ** **
        cells = [re.sub(r'\*\*(.+?)\*\*', r'\1', c) for c in cells]
        rows.append(cells)
        i += 1
    return header, rows, i

def render_inline(text):
    """Strip md inline syntax for plain text rendering."""
    text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
    text = re.sub(r'\*(.+?)\*', r'\1', text)
    text = re.sub(r'`(.+?)`', r'\1', text)
    return text

def add_inline_para(doc, text, size=12, italic=False):
    """Render markdown paragraph with bold spans preserved."""
    p = doc.add_paragraph()
    set_rtl_paragraph(p)
    p.paragraph_format.space_after = Pt(4)
    parts = re.split(r'(\*\*.+?\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = p.add_run(part[2:-2])
            style_run(run, size=size, bold=True, italic=italic)
        else:
            run = p.add_run(part)
            style_run(run, size=size, bold=False, italic=italic)

BLOCK_START = re.compile(
    r'^\s*(?:#{1,6}\s|\||[-*]\s|>|```|---\s*$|\d+(?:\.\d+)*\.\s|\*[^*].*\*\s*$)')

# Blocks that are complete in themselves - nothing may be folded into them.
# Bullets, numbered clauses and plain paragraphs do accept a folded line.
SELF_CLOSING = re.compile(r'^\s*(?:#{1,6}\s|\||---\s*$|\*[^*].*\*\s*$)')


def join_soft_wraps(lines):
    """Merge soft-wrapped lines into the block they belong to.

    In Markdown a single newline inside a paragraph is a soft wrap, not a line
    break - but the writer below emits one paragraph per source line, so a
    hard-wrapped source produced a paragraph break mid-sentence. Continuation
    lines are folded into the preceding block here; a blank line, a heading, a
    table row, a list item or a rule still starts a new one.
    """
    out = []
    for line in lines:
        stripped = line.strip()
        previous = out[-1].strip() if out else ''
        folds = (
            stripped                              # not a blank line
            and previous                          # something to fold into
            and not BLOCK_START.match(line)       # does not open a block itself
            and not SELF_CLOSING.match(out[-1])   # target accepts a fold
        )
        if folds:
            out[-1] = out[-1].rstrip() + ' ' + stripped
        else:
            out.append(line)
    return out
